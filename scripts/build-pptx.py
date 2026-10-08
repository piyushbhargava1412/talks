"""Turn Slidev's per-click PPTX export into an animated deck.

Starts from Slidev's own `--with-clicks` PPTX (one image slide per click
step, notes included — its notes parts open cleanly in Keynote, unlike
python-pptx's). For each Slidev slide, keeps the first step's slide and
stacks every later step on top as a full-slide picture that fades in on
click (native entrance animation), then drops the now-redundant step
slides. Every slide gets a push transition moving left (Slidev's
slide-left).

Normally run via scripts/export-pptx.sh, which does both exports.

usage: uv run --with python-pptx scripts/build-pptx.py <png_dir> <raw_slidev.pptx> <out.pptx> [--no-notes]
  png_dir: Slidev `--format png --with-clicks` export, named NNN-CC.png;
           only its names are used, for the step -> slide grouping. The
           images come from the PPTX itself so stacked steps line up exactly.
  --no-notes: leave the speaker notes out of the deck.
"""
import io
import sys
from collections import OrderedDict
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

args = [a for a in sys.argv[1:] if a != "--no-notes"]
keep_notes = len(args) == len(sys.argv) - 1
png_dir, raw_path, out_path = Path(args[0]), args[1], args[2]

prs = Presentation(raw_path)
pngs = sorted(png_dir.glob("*.png"))
slides = list(prs.slides)
assert len(pngs) == len(slides), (len(pngs), len(slides))

groups = OrderedDict()  # slide no -> [(png, pptx slide index)]
for i, png in enumerate(pngs):
    groups.setdefault(png.stem.split("-")[0], []).append((png, i))

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = f'xmlns:p="{P}"'
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def step_image(slide):
    """The step's screenshot, which Slidev sets as the slide background."""
    rid = slide._element.find(f".//{{{A}}}blip").get(f"{{{R}}}embed")
    return io.BytesIO(slide.part.related_part(rid).blob)


def click_fade(spid, ids):
    a, b, c, d, e = ids
    return f"""<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{c}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{d}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{e}" dur="300"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"""


drop = []
for steps in groups.values():
    slide = slides[steps[0][1]]
    sld = slide._element
    for tag in ("transition", "timing"):
        for old in sld.findall(f"{{{P}}}{tag}"):
            sld.remove(old)

    spids = [
        slide.shapes.add_picture(step_image(slides[i]), 0, 0, prs.slide_width, prs.slide_height).shape_id
        for _, i in steps[1:]
    ]
    drop += [i for _, i in steps[1:]]

    sld.append(etree.fromstring(f'<p:transition {NS} spd="med"><p:push dir="l"/></p:transition>'))
    if spids:
        next_id = iter(range(3, 10_000))
        clicks = "".join(click_fade(s, [next(next_id) for _ in range(5)]) for s in spids)
        sld.append(etree.fromstring(f"""<p:timing {NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks}</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""))

# Drop the per-step slides; unreferenced parts are not written on save.
sld_ids = prs.slides._sldIdLst
for i in sorted(drop, reverse=True):
    sld_id = sld_ids[i]
    prs.part.drop_rel(sld_id.rId)
    sld_ids.remove(sld_id)

# Unhook each slide from its notes page; like the dropped slides, the
# orphaned notes parts are not written on save.
if not keep_notes:
    for slide in prs.slides:
        for rid, rel in list(slide.part.rels.items()):
            if rel.reltype == RT.NOTES_SLIDE:
                slide.part.drop_rel(rid)

prs.save(out_path)
notes = "with" if keep_notes else "without"
print(f"{len(groups)} slides, {len(pngs)} click steps, {notes} speaker notes -> {out_path}")
